"""CSV and Excel exports. Amounts are written as decimal rupees (exact, from minor units)."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable

from app.core.money import from_minor


def _cell(v):
    return v


def to_csv(columns: list[str], rows: Iterable[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(columns)
    for r in rows:
        # Neutralise spreadsheet formula injection in text cells
        w.writerow([("'" + c) if isinstance(c, str) and c[:1] in ("=", "+", "-", "@") and not _is_number(c) else c for c in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def to_xlsx(sheets: dict[str, tuple[list[str], list[list]]]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title=name[:31])
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([("'" + c) if isinstance(c, str) and c[:1] == "=" else c for c in r])
        for i, col in enumerate(columns, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, min(60, len(col) + 4))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rupees(minor: int | None):
    return None if minor is None else from_minor(minor)


def dicts_to_table(rows: list[dict], money_keys: set[str] | None = None) -> tuple[list[str], list[list]]:
    if not rows:
        return ["(no data)"], []
    cols = list(rows[0].keys())
    money_keys = money_keys or {c for c in cols if c.endswith("_minor")}
    header = [c.replace("_minor", "") if c in money_keys else c for c in cols]
    out = []
    for r in rows:
        out.append([rupees(r[c]) if c in money_keys and isinstance(r[c], int) else
                    (", ".join(map(str, r[c])) if isinstance(r[c], list) else r[c]) for c in cols])
    return header, out
