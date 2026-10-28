"""File validation and content extraction. Everything runs locally; nothing leaves the machine."""

from __future__ import annotations

import csv
import io
import shutil
import zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

from app.core.config import get_settings
from app.services.imports.candidates import ExtractedDocument, ExtractedPage, ExtractedTable, Word


class ExtractionError(Exception):
    def __init__(self, message: str, code: str = "extraction_failed"):
        super().__init__(message)
        self.code = code


ALLOWED_EXTENSIONS = {".csv", ".txt", ".xlsx", ".xlsm", ".xls", ".pdf"}


def sniff_file_type(data: bytes, filename: str) -> str:
    """Identify the real file type from content, not just the extension."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError(f"Unsupported file extension {ext!r}", "unsupported_type")
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                names = z.namelist()
        except zipfile.BadZipFile as e:
            raise ExtractionError("Corrupt spreadsheet archive", "corrupt_file") from e
        if any(n.startswith("xl/") for n in names):
            return "xlsx"
        raise ExtractionError("ZIP file is not an Excel workbook", "unsupported_type")
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    head = data[:4096].lstrip().lower()
    if head.startswith(b"<") and (b"<table" in data[:200000].lower() or b"<html" in head):
        return "html_table"  # several banks export HTML with an .xls extension
    if b"\x00" in data[:8192]:
        raise ExtractionError("Binary file is not a supported statement format", "unsupported_type")
    return "csv"


def _decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode text file")  # pragma: no cover - latin-1 always decodes


def extract_csv(data: bytes, filename: str) -> ExtractedDocument:
    text = _decode_text(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    doc = ExtractedDocument(filename=filename, file_type="csv", text=text[:200000])
    doc.tables.append(ExtractedTable(rows=rows, source="csv"))
    return doc


def extract_xlsx(data: bytes, filename: str) -> ExtractedDocument:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ExtractionError(f"Could not open Excel workbook: {e.__class__.__name__}", "corrupt_file") from e
    doc = ExtractedDocument(filename=filename, file_type="xlsx")
    texts = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        doc.tables.append(ExtractedTable(rows=rows, source=f"sheet:{ws.title}"))
        texts.extend(" ".join(str(c) for c in r if c is not None) for r in rows[:200])
    wb.close()
    doc.text = "\n".join(texts)
    return doc


def extract_xls(data: bytes, filename: str) -> ExtractedDocument:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        raise ExtractionError(f"Could not open XLS workbook: {e.__class__.__name__}", "corrupt_file") from e
    doc = ExtractedDocument(filename=filename, file_type="xls")
    texts = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        doc.tables.append(ExtractedTable(rows=rows, source=f"sheet:{sheet.name}"))
        texts.extend(" ".join(str(c) for c in r if c not in (None, "")) for r in rows[:200])
    doc.text = "\n".join(texts)
    return doc


class _HTMLTables(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None
        self._depth = 0

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._depth += 1
            self.tables.append([])
        elif tag == "tr" and self._depth:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self.tables:
                self.tables[-1].append(self._row)
            self._row = None
        elif tag == "table":
            self._depth = max(0, self._depth - 1)

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def extract_html_table(data: bytes, filename: str) -> ExtractedDocument:
    p = _HTMLTables()
    p.feed(_decode_text(data))
    doc = ExtractedDocument(filename=filename, file_type="xls")
    doc.warnings.append("File has an .xls extension but contains an HTML table; read as HTML.")
    for i, t in enumerate(p.tables):
        if t:
            doc.tables.append(ExtractedTable(rows=t, source=f"html:table:{i + 1}"))
    doc.text = "\n".join(" ".join(r) for t in p.tables for r in t[:200])
    return doc


def tesseract_available() -> bool:
    s = get_settings()
    if not s.ocr_enabled:
        return False
    try:
        import pytesseract  # noqa: F401
    except ImportError:
        return False
    if s.tesseract_cmd:
        return Path(s.tesseract_cmd).exists()
    return shutil.which("tesseract") is not None


def _ocr_page(page, number: int) -> ExtractedPage:
    import pytesseract

    s = get_settings()
    if s.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = s.tesseract_cmd
    resolution = 300
    img = page.to_image(resolution=resolution).original
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT, config="--psm 6")
    scale = 72 / resolution
    words = []
    for i, txt in enumerate(data["text"]):
        txt = (txt or "").strip()
        if not txt or float(data["conf"][i]) < 0:
            continue
        x, y, w, h = (data[k][i] * scale for k in ("left", "top", "width", "height"))
        words.append(Word(text=txt, x0=x, x1=x + w, top=y, bottom=y + h))
    text = "\n".join(_words_to_lines_text(words))
    return ExtractedPage(number=number, text=text, words=words, width=float(page.width), ocr=True)


def _words_to_lines_text(words: list[Word]) -> list[str]:
    from app.services.imports.layout import group_lines

    return [" ".join(w.text for w in line) for line in group_lines(words)]


def extract_pdf(path: Path, filename: str, password: str | None = None, force_ocr: bool = False) -> ExtractedDocument:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    try:
        pdf = pdfplumber.open(str(path), password=password or "")
    except PDFPasswordIncorrect as e:
        if password:
            raise ExtractionError("The PDF password is incorrect.", "pdf_password_incorrect") from e
        raise ExtractionError("This PDF is password-protected. Enter its password to import it.",
                              "pdf_password_required") from e
    except Exception as e:
        raise ExtractionError(f"Could not open PDF: {e.__class__.__name__}", "corrupt_file") from e

    doc = ExtractedDocument(filename=filename, file_type="pdf", path=path)
    try:
        text_pages = []
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            raw_words = page.extract_words(x_tolerance=1.5, y_tolerance=2, keep_blank_chars=False)
            words = [Word(w["text"], w["x0"], w["x1"], w["top"], w["bottom"]) for w in raw_words]
            doc.pages.append(ExtractedPage(number=i, text=text, words=words, width=float(page.width)))
            text_pages.append(text)
            try:
                for j, t in enumerate(page.extract_tables(), start=1):
                    if t and len(t) > 1:
                        doc.tables.append(ExtractedTable(rows=t, source=f"page:{i}:table:{j}", page=i))
            except Exception:
                doc.warnings.append(f"Table detection failed on page {i}; text layout will be used.")
        total_chars = sum(len(t.strip()) for t in text_pages)
        doc.is_scanned = total_chars < 30 * max(1, len(pdf.pages))
        if doc.is_scanned or force_ocr:
            if tesseract_available():
                doc.pages = [_ocr_page(p, i) for i, p in enumerate(pdf.pages, start=1)]
                doc.tables = []
                doc.ocr_used = True
                doc.warnings.append(
                    "Scanned PDF: text was recognised with local OCR (Tesseract). "
                    "OCR can misread digits — review every amount and date."
                )
            else:
                doc.warnings.append(
                    "This PDF appears to be scanned (no text layer) and local OCR is not available. "
                    "Install Tesseract OCR (see docs/SETUP.md) or convert the statement to CSV/Excel."
                )
        doc.text = "\n".join(p.text for p in doc.pages)
    finally:
        pdf.close()
    return doc


def extract(data: bytes, filename: str, path: Path | None = None, password: str | None = None,
            force_ocr: bool = False) -> ExtractedDocument:
    kind = sniff_file_type(data, filename)
    if kind == "csv":
        return extract_csv(data, filename)
    if kind == "xlsx":
        return extract_xlsx(data, filename)
    if kind == "xls":
        return extract_xls(data, filename)
    if kind == "html_table":
        return extract_html_table(data, filename)
    if kind == "pdf":
        if path is None:
            raise ExtractionError("PDF extraction requires a stored file path")
        return extract_pdf(path, filename, password=password, force_ocr=force_ocr)
    raise ExtractionError(f"Unsupported file type {kind}", "unsupported_type")  # pragma: no cover
